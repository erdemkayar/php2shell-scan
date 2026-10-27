"""Report generation: HTML (self-contained), XLSX, DOCX, CSV, JSON. TR/EN.
Banner/header strings come from untrusted targets: HTML is escaped, spreadsheet cells are neutralised
against formula injection."""
import csv
import html
import json
import os
import re

from . import classify as C
from . import export as X
from . import ingest as IN
from . import __version__


T = {
    "en": {
        "title": "PHP2Shell exposure report", "sub": "CVE-2026-109618 / CVE-2026-109619 - PHP core RCE (fixed 2026-10-21)",
        "exec": "Executive summary", "assets": "Assets with live web services", "actionable": "Need action",
        "verified": "Host-verified", "patched": "Patched / no PHP", "by_status": "Distribution by status",
        "priority": "Priority actions", "all": "All assets", "search": "Filter...", "allst": "All statuses",
        "host": "Asset", "status": "Status", "conf": "Confidence", "php": "PHP version", "ports": "Ports",
        "group": "Group", "action": "Recommended action", "changes": "Changes since baseline",
        "new": "New assets", "gone": "No longer seen", "worse": "Newly exposed / worsened", "better": "Fixed / improved",
        "none": "None", "method": "Method and limitations", "remed": "Remediation", "interim": "Interim mitigations (reduce exposure, not a fix)",
        "sources": "Sources", "scope": "Scope", "endpoints": "Endpoints probed", "open": "Open services", "duration": "Duration",
        "generated": "Generated", "baseline": "Baseline", "groups": "By group", "total": "Total", "cms": "CMS hint",
        "notes": "Notes", "verified_by": "Verification", "closure": "Closure", "by_closure": "Remediation closure",
        "closure_line": "Closure: {c} of {p} PHP assets closed and verified on the host (patched, running the patched code, no IoC).",
        "compromise_line": " {s} asset(s) show compromise indicators: start incident response for these first.",
        "closure_changes": "Closure changes", "disagree": "Disagreements with external tools",
        "disagree_note": "An external tool (--targets-from) flagged these assets; this tool found them patched or without PHP on host, behavioural or banner evidence. Review which evidence is stronger.",
        "external": "External tool", "due": "Due", "owner": "Owner",
        "overdue_line": " {o} asset(s) are past their remediation deadline (oldest: {d} days).",
        "incomplete": "INCOMPLETE SCAN: {m} endpoint(s) were not scanned. Results are partial and must not be reported as complete.",
        "unknown_complete": "COMPLETENESS UNKNOWN: no scan metadata (meta.json) was found; treat the results as possibly partial.",
        "gaps": "COVERAGE GAPS: {n} target(s) were not scanned: {list}.",
        "scope_check": "Scope check", "advisory": "Advisory",
        "truncated": "Showing the first {n} of {total} (most urgent first). The complete list is in report.xlsx / assets.csv.",
        "headline_none": "No asset in scope requires action based on the evidence collected.",
        "headline": "{a} of {n} assets with live web services need action ({v} confirmed on the host). Patch and restart PHP workers first on the items below.",
        "method_items": [
            "Passive triage: ordinary HTTP(S) GET requests and response fingerprinting (X-Powered-By, Server, cookies, CMS markers). No exploit, memory-exhaustion or multipart trigger traffic is sent.",
            "TCP pre-filter on the listed ports, then HTTP(S) only on open ports. Requests are rate limited.",
            "Disclosed PHP versions are compared with the fixed releases 8.2.35, 8.3.36, 8.4.27 and 8.5.12. PHP 8.1 and older receive no routine upstream fix.",
            "Distribution builds (Ubuntu/Debian/RHEL) can carry the fix without changing the version string; such assets are marked 'verify' rather than vulnerable.",
            "'Verified' statuses come from read-only SSH checks on the host (php -v, package changelog for the CVE ids, running-process vs binary state).",
            "Limitations: stripped headers or a reverse proxy hide the real version; absence of a PHP indicator is not proof of absence; PHP used only via CLI is outside this scan.",
        ],
        "method_oracle": [
            "Active behavioural check (--oracle) was ENABLED: besides the GET requests, up to two ~400-byte multipart POSTs per web endpoint (invalid session id, upload-progress field, one-byte file). No exploit and no memory pressure.",
            "Pre-fix PHP prints a session-write warning that the fix removed: 'Vulnerable/Patched (behavioural check)' describe the RUNNING process, so they also catch distro backports and workers not restarted. They require the server to display PHP warnings; a silent answer is not evidence either way.",
        ],
        "remed_items": [
            "Upgrade to PHP 8.2.35 / 8.3.36 / 8.4.27 / 8.5.12 or later, or to a distribution package whose changelog lists CVE-2026-109618 and CVE-2026-109619.",
            "Restart PHP-FPM / Apache (mod_php) / FrankenPHP after the update; running workers keep the old code until restarted. "
            "Apache mod_php needs a full stop/start (systemctl restart apache2, not reload): in the lab, apachectl graceful and -k restart kept the old libphp loaded.",
            "PHP 8.1 and older: migrate to a supported branch or obtain a vendor backport.",
            "Hide version disclosure (expose_php=Off, remove X-Powered-By) after patching to reduce passive fingerprinting.",
        ],
        "interim_items": [
            "Limit request size and multipart part counts at the frontend; block multipart/form-data where the application does not need it (this breaks file uploads).",
            "session.upload_progress.enabled=0 at php.ini / pool level (vendors disagree on how much this helps; it is not sufficient alone).",
            "PHP 8.5: fatal_error_backtraces=0 and restart workers.",
            "Do not rely on RLIMIT/cgroup limits or on hiding errors; they do not prevent the issue.",
        ],
    },
    "tr": {
        "title": "PHP2Shell maruziyet raporu", "sub": "CVE-2026-109618 / CVE-2026-109619 - PHP core RCE (21.10.2026'da yamalandı)",
        "exec": "Yönetici özeti", "assets": "Canlı web servisi olan varlık", "actionable": "Aksiyon gereken",
        "verified": "Host'ta doğrulanan", "patched": "Yamalı / PHP yok", "by_status": "Duruma göre dağılım",
        "priority": "Öncelikli aksiyonlar", "all": "Tüm varlıklar", "search": "Filtrele...", "allst": "Tüm durumlar",
        "host": "Varlık", "status": "Durum", "conf": "Güven", "php": "PHP sürümü", "ports": "Portlar",
        "group": "Grup", "action": "Önerilen aksiyon", "changes": "Baseline'a göre değişiklikler",
        "new": "Yeni varlıklar", "gone": "Artık görülmeyen", "worse": "Yeni maruz / kötüleşen", "better": "Düzelen / iyileşen",
        "none": "Yok", "method": "Yöntem ve sınırlamalar", "remed": "Düzeltme", "interim": "Geçici önlemler (maruziyeti azaltır, çözüm değildir)",
        "sources": "Kaynaklar", "scope": "Kapsam", "endpoints": "Taranan endpoint", "open": "Açık servis", "duration": "Süre",
        "generated": "Oluşturulma", "baseline": "Baseline", "groups": "Gruplara göre", "total": "Toplam", "cms": "CMS ipucu",
        "notes": "Notlar", "verified_by": "Doğrulama", "closure": "Kapanış", "by_closure": "Düzeltme kapanışı",
        "closure_line": "Kapanış: {p} PHP varlığının {c} tanesi host üzerinde kanıtlı kapandı (yamalı, çalışan süreç yamalı, IoC yok).",
        "compromise_line": " {s} varlıkta ele geçirilme izi var: önce bunlar için olay müdahalesi başlatın.",
        "closure_changes": "Kapanış değişiklikleri", "disagree": "Dış araçlarla uyuşmazlıklar",
        "disagree_note": "Dış bir araç (--targets-from) bu varlıklar için bulgu verdi; bu araç host, davranış veya banner kanıtıyla yamalı ya da PHP'siz buldu. Hangi kanıtın daha güçlü olduğunu inceleyin.",
        "external": "Dış araç", "due": "Son tarih", "owner": "Sorumlu",
        "overdue_line": " {o} varlığın düzeltme süresi doldu (en eski: {d} gün gecikme).",
        "incomplete": "EKSİK TARAMA: {m} endpoint taranmadı. Sonuçlar kısmidir, tamamlanmış tarama olarak raporlanmamalıdır.",
        "unknown_complete": "TAMAMLANMA BİLİNMİYOR: tarama meta verisi (meta.json) bulunamadı; sonuçları kısmi olabilir kabul edin.",
        "gaps": "KAPSAMA BOŞLUĞU: {n} hedef taranmadı: {list}.",
        "scope_check": "Kapsam kontrolü", "advisory": "Advisory",
        "truncated": "İlk {n} / {total} gösteriliyor (en acil olanlar önce). Tam liste report.xlsx / assets.csv içinde.",
        "headline_none": "Toplanan kanıtlara göre kapsamdaki hiçbir varlık aksiyon gerektirmiyor.",
        "headline": "Canlı web servisi olan {n} varlığın {a} tanesi aksiyon gerektiriyor ({v} tanesi host üzerinde doğrulandı). Önce aşağıdaki kalemlerde PHP'yi yamalayıp worker'ları yeniden başlatın.",
        "method_items": [
            "Pasif triage: standart HTTP(S) GET istekleri ve yanıt parmak izi (X-Powered-By, Server, cookie, CMS işaretleri). Exploit, bellek tüketimi veya multipart tetikleme trafiği gönderilmez.",
            "Listelenen portlarda TCP ön eleme, ardından yalnızca açık portlara HTTP(S). İstekler hız sınırlıdır.",
            "Açıklanan PHP sürümleri 8.2.35, 8.3.36, 8.4.27 ve 8.5.12 düzeltilmiş sürümleriyle karşılaştırılır. PHP 8.1 ve altı için rutin upstream düzeltme yoktur.",
            "Dağıtım build'leri (Ubuntu/Debian/RHEL) sürüm dizesini değiştirmeden yamayı taşıyabilir; bu varlıklar 'savunmasız' değil 'doğrulanmalı' olarak işaretlenir.",
            "'Doğrulanmış' durumlar host üzerindeki salt okunur SSH kontrollerinden gelir (php -v, CVE numaraları için paket changelog'u, çalışan süreç ile binary durumu).",
            "Sınırlamalar: temizlenmiş header veya reverse proxy gerçek sürümü gizler; PHP göstergesinin olmaması yokluğun kanıtı değildir; yalnızca CLI'dan kullanılan PHP bu tarama dışındadır.",
        ],
        "method_oracle": [
            "Aktif davranış kontrolü (--oracle) AÇIKTI: GET isteklerine ek olarak web endpoint'i başına en fazla iki adet ~400 baytlık multipart POST (geçersiz session id, upload progress alanı, tek baytlık dosya). Exploit ve bellek baskısı yoktur.",
            "Yama öncesi PHP, yamanın kaldırdığı session-write uyarısını basar: 'Savunmasız/Yamalı (davranış kontrolü)' durumları ÇALIŞAN süreci anlatır; distro backport'unu ve yeniden başlatılmamış worker'ı da yakalar. Sunucunun PHP uyarılarını göstermesi gerekir; sessiz yanıt iki yönde de kanıt değildir.",
        ],
        "remed_items": [
            "PHP 8.2.35 / 8.3.36 / 8.4.27 / 8.5.12 veya üstüne ya da changelog'unda CVE-2026-109618 ve CVE-2026-109619 yer alan dağıtım paketine yükseltin.",
            "Güncellemeden sonra PHP-FPM / Apache (mod_php) / FrankenPHP'yi yeniden başlatın; çalışan worker'lar yeniden başlatılana kadar eski kodu kullanır. "
            "Apache mod_php tam stop/start ister (systemctl restart apache2, reload değil): lab'da apachectl graceful ve -k restart eski libphp'yi yüklü tuttu.",
            "PHP 8.1 ve altı: desteklenen bir branch'e geçin veya vendor backport'u alın.",
            "Yamadan sonra sürüm ifşasını kapatın (expose_php=Off, X-Powered-By'ı kaldırın); pasif parmak izini azaltır.",
        ],
        "interim_items": [
            "Frontend'de istek boyutunu ve multipart part sayısını sınırlayın; uygulama gerektirmiyorsa multipart/form-data'yı engelleyin (dosya yüklemeyi bozar).",
            "php.ini / pool seviyesinde session.upload_progress.enabled=0 (vendor'lar faydası konusunda hemfikir değil; tek başına yeterli değil).",
            "PHP 8.5: fatal_error_backtraces=0 ve worker'ları yeniden başlatın.",
            "RLIMIT/cgroup limitlerine veya hata gizlemeye güvenmeyin; sorunu engellemezler.",
        ],
    },
}

HTML_MAX_ROWS, HTML_MAX_PRIO = 2000, 1000
COLORS = {"critical": "#b42318", "high": "#d92d20", "medium": "#dc6803", "ok": "#12803c", "info": "#667085"}


def _incomplete(meta, lang):
    """Coverage warning shown at the top of every report ('' when the scan covered everything it was given)."""
    t, parts = T[lang], []
    if meta.get("complete") is None:
        parts.append(t["unknown_complete"])
    elif meta.get("complete") is False:
        parts.append(t["incomplete"].format(m=meta.get("missing", 0)))
    gaps = meta.get("gaps") or []
    if gaps:
        shown = ", ".join(gaps[:10]) + (f" (+{len(gaps) - 10})" if len(gaps) > 10 else "")
        parts.append(t["gaps"].format(n=len(gaps), list=shown))
    return " ".join(parts)


def _texts(lang, meta):
    """Report strings for this run: the method section must say whether --oracle sent POSTs; an advisory may rename
    the report ("texts": {"en": {"title": ..., "sub": ...}})."""
    over = {k: v for k, v in (C.ADVISORY.get("texts", {}).get(lang) or {}).items() if k in ("title", "sub")}
    t = dict(T[lang], **over)
    if not meta.get("oracle"):
        return t
    return dict(t, method_items=t["method_oracle"] + t["method_items"][1:])


def _meta_lines(meta):
    return [("advisory", f"{C.ADVISORY.get('id', '')} {C.ADVISORY.get('published', '')}".strip()), ("scope", meta.get("scope", "")), ("scope_check", meta.get("scope_check", "")), ("endpoints", f"{meta.get('endpoints', 0)}"),
            ("open", f"{meta.get('open', 0)}"), ("duration", meta.get("duration", ""))]


def _exec_text(summary, lang):
    t = T[lang]
    head = t["headline_none"] if not summary["actionable"] else t["headline"].format(
        a=summary["actionable"], n=summary["assets"],
        v=sum(summary["by_status"].get(k, 0) for k in ("VERIFIED_VULNERABLE", "VERIFIED_RESTART_REQUIRED")))
    cl = summary.get("by_closure") or {}
    php = sum(v for k, v in cl.items() if k not in ("NO_PHP_EVIDENCE", "NOT_APPLICABLE"))
    if php:
        head += " " + t["closure_line"].format(c=cl.get("CLOSED_VERIFIED", 0), p=php)
    if cl.get("COMPROMISE_SUSPECTED"):
        head += t["compromise_line"].format(s=cl["COMPROMISE_SUSPECTED"])
    if summary.get("overdue"):
        head += t["overdue_line"].format(o=summary["overdue"], d=summary.get("max_overdue_days", 0))
    return head


def _esc(x):
    return html.escape("" if x is None else str(x), quote=True)


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def safe_cell(v):
    """Neutralise spreadsheet formula injection and strip control characters (they abort XLSX/DOCX writers)."""
    if isinstance(v, str):
        v = _CTRL.sub("", v)
        if v[:1] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + v
    return v


def _asset_row(a, lang):
    return {
        "host": a["host"], "names": ", ".join(a["names"]), "status": a["final_status"],
        "status_label": C.status_label(a["final_status"], lang), "confidence": a["confidence"],
        "php": ", ".join(a["php_versions"]), "ports": ",".join(map(str, a["ports"])),
        "cms": ", ".join(a["cms"]), "server": ", ".join(a["servers"]), "group": a["labels"].get("customer", "-"),
        "env": a["labels"].get("env", ""), "action": C.status_action(a["final_status"], lang), "note": a["note"],
        "verified": (a["verified"] or {}).get("summary", ""),
        "closure": a.get("closure", ""), "closure_label": C.closure_label(a.get("closure", ""), lang),
        "needs_action": C.is_actionable(a["final_status"]) or a.get("closure") == "COMPROMISE_SUSPECTED",
        "due": (a.get("sla") or {}).get("due", ""), "overdue_days": (a.get("sla") or {}).get("overdue_days", ""),
        "owner": (a.get("sla") or {}).get("owner") or a["labels"].get("owner", ""),
    }


# ----------------------------------------------------------------------------- HTML
def html_report(assets, summary, groups, meta, diff, lang="en"):
    t = _texts(lang, meta)
    # HTML is for reading: cap the tables (assets are sorted by urgency) and point to XLSX/CSV for the full list
    all_rows = [_asset_row(a, lang) for a in assets]
    rows = all_rows[:HTML_MAX_ROWS]
    prio_all = [r for r in all_rows if r["needs_action"]]
    n_rows, n_prio = len(assets), len(prio_all)
    prio = prio_all[:HTML_MAX_PRIO]
    more = lambda shown, total: (f'<p class="mut">{_esc(t["truncated"].format(n=shown, total=total))}</p>' if total > shown else "")
    order = [s for s in sorted(C.STATUS, key=C.rank) if summary["by_status"].get(s)]
    mx = max([summary["by_status"].get(s, 0) for s in order] + [1])

    def badge(s):
        return f'<span class="b" style="--c:{COLORS[C.severity(s)]}">{_esc(C.status_label(s, lang))}</span>'

    def cbadge(c):
        return f'<span class="b" style="--c:{COLORS[C.CLOSURE.get(c, ("info",))[0]]}">{_esc(C.closure_label(c, lang))}</span>'

    by_cl = summary.get("by_closure") or {}
    corder = [c for c in C.CLOSURE if by_cl.get(c)]
    cbars = "".join(f'<div class="bar"><span class="bl">{cbadge(c)}</span><span class="bt"><i style="width:{100*by_cl[c]/max(by_cl.values()):.1f}%;'
                    f'background:{COLORS[C.CLOSURE[c][0]]}"></i></span><b>{by_cl[c]}</b></div>' for c in corder)

    bars = "".join(
        f'<div class="bar"><span class="bl">{badge(s)}</span><span class="bt"><i style="width:{100*summary["by_status"][s]/mx:.1f}%;background:{COLORS[C.severity(s)]}"></i></span>'
        f'<b>{summary["by_status"][s]}</b></div>' for s in order)

    def table(rs, cols, cls=""):
        th = "".join(f"<th>{_esc(t[c] if c in t else c)}</th>" for c, _ in cols)
        body = ""
        for r in rs:
            tds = ""
            for _, f in cols:
                tds += f"<td>{f(r)}</td>"
            body += f'<tr data-s="{_esc(r.get("status", ""))}">{tds}</tr>'
        return f'<table class="{cls}"><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table>'

    prio_cols = [("host", lambda r: f"<b>{_esc(r['host'])}</b><br><small>{_esc(r['names'])}</small>"),
                 ("group", lambda r: _esc(r["group"])), ("status", lambda r: badge(r["status"])),
                 ("closure", lambda r: cbadge(r["closure"])), ("php", lambda r: _esc(r["php"] or "-")),
                 ("ports", lambda r: _esc(r["ports"])),
                 ("due", lambda r: _esc(r["due"]) + (f' <b style="color:#d92d20">+{r["overdue_days"]}</b>' if r["overdue_days"] else "")),
                 ("owner", lambda r: _esc(r["owner"])), ("action", lambda r: _esc(r["action"]))]
    all_cols = [("host", lambda r: f"<b>{_esc(r['host'])}</b><br><small>{_esc(r['names'])}</small>"),
                ("group", lambda r: _esc(r["group"])), ("status", lambda r: badge(r["status"])),
                ("closure", lambda r: cbadge(r["closure"])),
                ("conf", lambda r: _esc(r["confidence"])), ("php", lambda r: _esc(r["php"] or "-")),
                ("ports", lambda r: _esc(r["ports"])), ("cms", lambda r: _esc(r["cms"])),
                ("notes", lambda r: _esc(r["verified"] or r["note"]))]
    grp_rows = [{"status": "", "group": g["group"], "total": g["total"], "actionable": g["actionable"]} for g in groups]
    grp_cols = [("group", lambda r: _esc(r["group"])), ("total", lambda r: str(r["total"])),
                ("actionable", lambda r: f"<b>{r['actionable']}</b>")]

    chg = ""
    if diff:
        def lst(title, items, f):
            li = "".join(f"<li>{f(i)}</li>" for i in items[:25]) or f"<li>{_esc(t['none'])}</li>"
            if len(items) > 25:
                li += f"<li class='mut'>+{len(items) - 25} ... (assets.json / XLSX)</li>"
            return f"<div><h3>{_esc(title)} ({len(items)})</h3><ul>{li}</ul></div>"
        fs = lambda i: f"{_esc(i['host'])} - {_esc(C.status_label(i.get('status') or i.get('to'), lang))}"
        ft = lambda i: f"{_esc(i['host'])}: {_esc(C.status_label(i['from'], lang))} → {_esc(C.status_label(i['to'], lang))}"
        fc = lambda i: f"{_esc(i['host'])}: {_esc(C.closure_label(i['from'], lang))} → {_esc(C.closure_label(i['to'], lang))}"
        chg = (f'<section><h2>{_esc(t["changes"])}</h2><p class="mut">{_esc(t["baseline"])}: {_esc(meta.get("baseline", ""))}</p>'
               f'<div class="cols">{lst(t["worse"], diff["worsened"], ft)}{lst(t["better"], diff["improved"], ft)}'
               f'{lst(t["new"], diff["new_assets"], fs)}{lst(t["gone"], diff["gone_assets"], fs)}'
               f'{lst(t["closure_changes"], diff.get("closure_changes", []), fc)}</div></section>')

    dis = IN.disagreements(assets)
    dis_html = "" if not dis else (
        f'<section><h2>{_esc(t["disagree"])} ({len(dis)})</h2><p class="mut">{_esc(t["disagree_note"])}</p><div class="wrap">'
        + table([dict(d, status=d["status"]) for d in dis],
                [("host", lambda r: _esc(r["host"])), ("external", lambda r: _esc(r["external"])),
                 ("status", lambda r: badge(r["status"])), ("closure", lambda r: cbadge(r["closure"]))]) + "</div></section>")
    kpi = [(t["assets"], summary["assets"], ""), (t["actionable"], summary["actionable"], "bad" if summary["actionable"] else "good"),
           (t["verified"], summary["verified"], ""),
           (t["patched"], sum(summary["by_status"].get(k, 0) for k in ("PATCHED_BY_BANNER", "ORACLE_PATCHED", "VERIFIED_PATCHED",
                                                                        "NO_PHP_INDICATOR", "VERIFIED_NO_PHP")), "good")]
    kpis = "".join(f'<div class="k {c}"><b>{v}</b><span>{_esc(n)}</span></div>' for n, v, c in kpi)
    ul = lambda xs: "<ul>" + "".join(f"<li>{_esc(x)}</li>" for x in xs) + "</ul>"
    opts = "".join(f'<option value="{_esc(s)}">{_esc(C.status_label(s, lang))}</option>' for s in order)
    src = "".join((f'<li><a href="{_esc(u)}" rel="noopener">{_esc(n)}</a></li>' if str(u).startswith(("https://", "http://"))
                   else f"<li>{_esc(n)} - {_esc(u)}</li>") for n, u in C.ADVISORY.get("sources", []))
    meta_bits = " · ".join(f"{_esc(t[k])}: {_esc(v)}" for k, v in _meta_lines(meta) if v != "")

    return f"""<!doctype html><html lang="{lang}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{_esc(t['title'])}</title><style>
:root{{--bg:#fff;--fg:#101828;--mut:#667085;--card:#f9fafb;--line:#e4e7ec;--acc:#1f3a5f}}
@media(prefers-color-scheme:dark){{:root{{--bg:#0c111d;--fg:#f2f4f7;--mut:#98a2b3;--card:#161b26;--line:#2a3140;--acc:#9bb8e8}}}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--fg);font:14px/1.5 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}}
main{{max-width:1180px;margin:0 auto;padding:24px 16px 64px}}h1{{margin:0;font-size:26px;color:var(--acc)}}h2{{margin:32px 0 10px;font-size:18px;border-bottom:1px solid var(--line);padding-bottom:6px}}
h3{{font-size:14px;margin:0 0 6px}}.mut,small{{color:var(--mut)}}.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;margin:16px 0}}
.k{{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px}}.k b{{display:block;font-size:28px}}.k.bad b{{color:#d92d20}}.k.good b{{color:#12803c}}.k span{{color:var(--mut)}}
.sum{{background:var(--card);border-left:4px solid var(--acc);padding:12px 14px;border-radius:6px}}
.bar{{display:grid;grid-template-columns:minmax(220px,320px) 1fr 40px;gap:10px;align-items:center;margin:6px 0}}.bt{{background:var(--line);height:10px;border-radius:5px;overflow:hidden}}.bt i{{display:block;height:100%}}
.b{{display:inline-block;padding:2px 8px;border-radius:10px;border:1px solid var(--c);color:var(--c);font-size:12px;font-weight:600;white-space:nowrap}}
table{{width:100%;border-collapse:collapse;font-size:13px}}th,td{{text-align:left;padding:8px 10px;border-bottom:1px solid var(--line);vertical-align:top}}th{{position:sticky;top:0;background:var(--card)}}
.tools{{display:flex;gap:8px;margin:8px 0}}input,select{{padding:6px 8px;border:1px solid var(--line);border-radius:6px;background:var(--bg);color:var(--fg)}}
.cols{{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:16px}}ul{{margin:4px 0 4px 18px;padding:0}}
.wrap{{overflow-x:auto}}.warn{{background:#fef3f2;color:#b42318;border:1px solid #fda29b;border-radius:8px;padding:12px 14px;font-weight:700}}@media print{{.tools{{display:none}}th{{position:static}}}}
</style></head><body><main>
<h1>{_esc(meta.get('title') or t['title'])}</h1><div class="mut">{_esc(t['sub'])}</div>
<div class="mut">{_esc(t['generated'])}: {_esc(meta.get('generated', ''))} · php2shell-scanner {__version__} · {meta_bits}</div>
{('<p class="warn">' + _esc(_incomplete(meta, lang)) + '</p>') if _incomplete(meta, lang) else ''}
<section><h2>{_esc(t['exec'])}</h2><div class="kpis">{kpis}</div><p class="sum">{_esc(_exec_text(summary, lang))}</p>
<h3>{_esc(t['by_status'])}</h3>{bars}<h3>{_esc(t['by_closure'])}</h3>{cbars}</section>
<section><h2>{_esc(t['priority'])} ({n_prio})</h2>{more(len(prio), n_prio)}<div class="wrap">{table(prio, prio_cols) if prio else '<p>'+_esc(t['none'])+'</p>'}</div></section>
<section><h2>{_esc(t['groups'])}</h2>{table(grp_rows, grp_cols)}</section>
{chg}
{dis_html}
<section><h2>{_esc(t['all'])} ({n_rows})</h2>{more(len(rows), n_rows)}<div class="tools"><input id="q" placeholder="{_esc(t['search'])}"><select id="st"><option value="">{_esc(t['allst'])}</option>{opts}</select></div>
<div class="wrap" id="all">{table(rows, all_cols)}</div></section>
<section><h2>{_esc(t['remed'])}</h2>{ul(t['remed_items'])}<h3>{_esc(t['interim'])}</h3>{ul(t['interim_items'])}</section>
<section><h2>{_esc(t['method'])}</h2>{ul(t['method_items'])}<h3>{_esc(t['sources'])}</h3><ul>{src}</ul>{ul(C.ADVISORY.get('evidence', []))}</section>
<script>
(function(){{var q=document.getElementById('q'),s=document.getElementById('st'),rows=document.querySelectorAll('#all tbody tr');
function f(){{var x=q.value.toLowerCase(),y=s.value;rows.forEach(function(r){{r.style.display=(r.textContent.toLowerCase().indexOf(x)>-1&&(!y||r.dataset.s===y))?'':'none'}})}}
q.addEventListener('input',f);s.addEventListener('change',f)}})();
</script></main></body></html>"""


# ----------------------------------------------------------------------------- XLSX
class _Xlsx:
    """Write-only (streaming) workbook helpers: memory stays flat for /16-sized scans."""

    def __init__(self, lang):
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font, PatternFill
        self.t, self.lang = T[lang], lang
        self.wb, self._Cell = Workbook(write_only=True), WriteOnlyCell
        self.HDR_FILL, self.HDR_FONT = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
        self.BOLD, self.WRAP = Font(bold=True), Alignment(vertical="top", wrap_text=True)
        self.FILL = {k: PatternFill("solid", fgColor=v) for k, v in
                     {"critical": "F4C7C3", "high": "FAD4D0", "medium": "FDE7C8", "ok": "D3F0DC", "info": "E4E7EC"}.items()}
        self.label_sev = {C.status_label(k, lang): C.severity(k) for k in C.STATUS}
        self.label_sev.update({k: C.severity(k) for k in C.STATUS})
        self.label_sev.update({C.closure_label(k, lang): v[0] for k, v in C.CLOSURE.items()})
        self.Font = Font

    def cell(self, ws, v, font=None, fill=None, align=None):
        c = self._Cell(ws, value=safe_cell(v))
        if font:
            c.font = font
        if fill:
            c.fill = fill
        if align:
            c.alignment = align
        return c

    def sheet(self, name, headers, rows, widths=None, status_col=None):
        from openpyxl.utils import get_column_letter
        ws = self.wb.create_sheet(name)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i) or min(max(len(str(h)) + 4, 14), 48)
        ws.freeze_panes = "A2"
        ws.append([self.cell(ws, h, self.HDR_FONT, self.HDR_FILL) for h in headers])
        n = 0
        for r in rows:
            ws.append([self.cell(ws, v, fill=self.FILL.get(self.label_sev.get(v)) if ci == status_col else None, align=self.WRAP)
                       for ci, v in enumerate(r)])
            n += 1
        if n:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n + 1}"
        return ws

    def summary_sheet(self, summary, meta, order):
        from openpyxl.chart import BarChart, Reference
        t, ws = self.t, self.wb.create_sheet("Summary")
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 46, 30
        ws.append([self.cell(ws, meta.get("title") or t["title"], self.Font(bold=True, size=16, color="1F3A5F"))])
        ws.append([t["sub"]])
        ws.append([f"{t['generated']}: {meta.get('generated', '')}  |  php2shell-scanner {__version__}  |  "
                   f"{C.ADVISORY.get('id', '')} ({', '.join(C.CVES)})"])
        warn = _incomplete(meta, self.lang)
        ws.append([self.cell(ws, warn, self.Font(bold=True, color="B42318"))] if warn else [])
        for n, v in [(t["assets"], summary["assets"]), (t["actionable"], summary["actionable"]),
                     (t["verified"], summary["verified"]), (t["endpoints"], meta.get("endpoints", 0)),
                     (t["open"], meta.get("open", 0)), (t["duration"], meta.get("duration", "")),
                     (t["scope"], meta.get("scope", "")), (t["scope_check"], meta.get("scope_check", ""))]:
            ws.append([self.cell(ws, n, self.BOLD), self.cell(ws, v)])
        ws.append([self.cell(ws, _exec_text(summary, self.lang), align=self.WRAP)])
        ws.append([])
        head_row = 4 + 8 + 2 + 1  # rows written so far + 1
        ws.append([self.cell(ws, t["status"], self.BOLD), self.cell(ws, t["total"], self.BOLD)])
        for st in order:
            ws.append([self.cell(ws, C.status_label(st, self.lang), fill=self.FILL[C.severity(st)]), summary["by_status"][st]])
        ws.append([])
        ws.append([self.cell(ws, t["by_closure"], self.BOLD), self.cell(ws, t["total"], self.BOLD)])
        for c in C.CLOSURE:
            if (summary.get("by_closure") or {}).get(c):
                ws.append([self.cell(ws, C.closure_label(c, self.lang), fill=self.FILL[C.CLOSURE[c][0]]), summary["by_closure"][c]])
        if order:
            ch = BarChart()
            ch.type, ch.title, ch.legend = "bar", t["by_status"], None
            ch.add_data(Reference(ws, min_col=2, min_row=head_row, max_row=head_row + len(order)), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=head_row + 1, max_row=head_row + len(order)))
            ch.height, ch.width, ch.anchor = 8, 18, "D5"
            ws.add_chart(ch)

    def method_sheet(self):
        t, wm = self.t, self.wb.create_sheet("Method")
        wm.column_dimensions["A"].width = 130
        for h, items in ((t["method"], t["method_items"]), (t["remed"], t["remed_items"]), (t["interim"], t["interim_items"])):
            wm.append([self.cell(wm, h, self.BOLD)])
            for it in items:
                wm.append([self.cell(wm, it, align=self.WRAP)])
            wm.append([])
        wm.append([self.cell(wm, t["sources"], self.BOLD)])
        for n, u in C.ADVISORY.get("sources", []):
            wm.append([f"{n} - {u}"])
        for e in C.ADVISORY.get("evidence", []):
            wm.append([self.cell(wm, e, align=self.WRAP)])


def _changes_rows(diff, t, lang):
    cr = []
    for k, lab in (("worsened", t["worse"]), ("improved", t["better"])):
        cr += [[lab, i["host"], C.status_label(i["from"], lang), C.status_label(i["to"], lang)] for i in diff[k]]
    cr += [[t["closure_changes"], i["host"], C.closure_label(i["from"], lang), C.closure_label(i["to"], lang)]
           for i in diff.get("closure_changes", [])]
    cr += [[t["new"], i["host"], "", C.status_label(i["status"], lang)] for i in diff["new_assets"]]
    cr += [[t["gone"], i["host"], C.status_label(i["status"], lang), ""] for i in diff["gone_assets"]]
    return cr


def xlsx_report(path, assets, summary, groups, meta, diff, endpoint_rows, lang="en"):
    x, t = _Xlsx(lang), _texts(lang, meta)
    x.t = t
    order = [s for s in sorted(C.STATUS, key=C.rank) if summary["by_status"].get(s)]
    x.summary_sheet(summary, meta, order)
    H = [t["host"], "SNI/SAN", t["group"], "env", t["status"], t["closure"], t["conf"], t["php"], t["ports"], t["cms"], "Server",
         t["action"], t["notes"], t["verified_by"], t["due"], "overdue_days", t["owner"]]

    def mk(a):
        r = _asset_row(a, lang)
        return [r["host"], r["names"], r["group"], r["env"], r["status_label"], r["closure_label"], r["confidence"], r["php"],
                r["ports"], r["cms"], r["server"], r["action"], r["note"], r["verified"], r["due"], r["overdue_days"], r["owner"]]
    W = {1: 20, 2: 28, 5: 34, 6: 34, 12: 60, 13: 50, 14: 50}
    x.sheet("Priority", H, (mk(a) for a in assets if _asset_row(a, lang)["needs_action"]), W, 4)
    x.sheet("Assets", H, (mk(a) for a in assets), W, 4)
    x.sheet("Endpoints", ["host", "port", "scheme", "sni", "status", "confidence", "php_version", "server", "x_powered_by",
                          "http_status", "indicators", "cms", "tls_cn", "tls_not_after", "note"], endpoint_rows, {15: 60}, 4)
    x.sheet("Groups", [t["group"], t["total"], t["actionable"]] + [C.status_label(s, lang) for s in order],
            ([g["group"], g["total"], g["actionable"]] + [g["by_status"].get(s, 0) for s in order] for g in groups))
    if diff:
        x.sheet("Changes", ["type", t["host"], "from", "to"], _changes_rows(diff, t, lang), {1: 30, 2: 24, 3: 36, 4: 36})
    x.method_sheet()
    x.wb.save(path)


# ----------------------------------------------------------------------------- DOCX
def docx_report(path, assets, summary, groups, meta, diff, lang="en"):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt, Cm

    t = _texts(lang, meta)
    d = Document()
    s = d.sections[0]
    s.orientation, s.page_width, s.page_height = WD_ORIENT.LANDSCAPE, s.page_height, s.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Cm(1.8))
    d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Calibri", Pt(10)
    d.add_heading(meta.get("title") or t["title"], 0)
    d.add_paragraph(f"{t['sub']}\n{t['generated']}: {meta.get('generated', '')} · php2shell-scanner {__version__}")
    if _incomplete(meta, lang):
        p = d.add_paragraph()
        r = p.add_run(_incomplete(meta, lang))
        r.bold = True
        from docx.shared import RGBColor
        r.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
    d.add_heading(t["exec"], 1)
    d.add_paragraph(_exec_text(summary, lang))
    tb = d.add_table(rows=1, cols=2); tb.style = "Light Grid Accent 1"
    tb.rows[0].cells[0].text, tb.rows[0].cells[1].text = t["status"], t["total"]
    for st in sorted(summary["by_status"], key=C.rank):
        c = tb.add_row().cells
        c[0].text, c[1].text = C.status_label(st, lang), str(summary["by_status"][st])
    d.add_heading(t["by_closure"], 2)
    tb = d.add_table(rows=1, cols=2); tb.style = "Light Grid Accent 1"
    tb.rows[0].cells[0].text, tb.rows[0].cells[1].text = t["closure"], t["total"]
    for cl in C.CLOSURE:
        if (summary.get("by_closure") or {}).get(cl):
            c = tb.add_row().cells
            c[0].text, c[1].text = C.closure_label(cl, lang), str(summary["by_closure"][cl])
    rows = [_asset_row(a, lang) for a in assets]
    prio = [r for r in rows if r["needs_action"]]
    d.add_heading(f"{t['priority']} ({len(prio)})", 1)
    if prio:
        tb = d.add_table(rows=1, cols=7); tb.style = "Light Grid Accent 1"
        for i, h in enumerate([t["host"], t["group"], t["status"], t["closure"], t["php"], t["ports"], t["action"]]):
            tb.rows[0].cells[i].text = h
        for r in prio[:200]:
            c = tb.add_row().cells
            for i, v in enumerate([r["host"], r["group"], r["status_label"], r["closure_label"], r["php"] or "-", r["ports"], r["action"]]):
                c[i].text = _CTRL.sub("", str(v))
        if len(prio) > 200:
            d.add_paragraph(f"... +{len(prio) - 200} (see XLSX)")
    else:
        d.add_paragraph(t["none"])
    d.add_heading(t["groups"], 1)
    tb = d.add_table(rows=1, cols=3); tb.style = "Light Grid Accent 1"
    for i, h in enumerate([t["group"], t["total"], t["actionable"]]):
        tb.rows[0].cells[i].text = h
    for g in groups:
        c = tb.add_row().cells
        c[0].text, c[1].text, c[2].text = _CTRL.sub("", str(g["group"])), str(g["total"]), str(g["actionable"])
    if diff:
        d.add_heading(t["changes"], 1)
        for k, lab in (("worsened", t["worse"]), ("improved", t["better"]), ("new_assets", t["new"]), ("gone_assets", t["gone"]),
                       ("closure_changes", t["closure_changes"])):
            d.add_paragraph(f"{lab}: {len(diff.get(k, []))}", style="List Bullet")
    for h, items in ((t["remed"], t["remed_items"]), (t["interim"], t["interim_items"]), (t["method"], t["method_items"])):
        d.add_heading(h, 1)
        for it in items:
            d.add_paragraph(it, style="List Bullet")
    d.add_heading(t["sources"], 1)
    for n, u in C.ADVISORY.get("sources", []):
        d.add_paragraph(f"{n} - {u}", style="List Bullet")
    d.save(path)


# ----------------------------------------------------------------------------- CSV/JSON + driver
def write_all(outdir, results, assets, summary, groups, meta, diff, formats, lang="en"):
    os.makedirs(outdir, exist_ok=True)
    written = []
    ep_rows = ([r["host"], r["port"], r.get("scheme", ""), r.get("sni", ""), r["status"], r.get("confidence", ""),
                r.get("php_version", ""), r.get("server", ""), r.get("x_powered_by", ""), r.get("http_status", ""),
                ";".join(r.get("indicators", [])), r.get("cms_hint", ""), (r.get("tls") or {}).get("cn", ""),
                (r.get("tls") or {}).get("not_after", ""), r.get("note", "")]
               for r in results if r["status"] not in ("CLOSED", "ERROR"))
    if "json" in formats:
        p = os.path.join(outdir, "assets.json")
        json.dump({"meta": meta, "summary": summary, "assets": assets, "groups": groups, "diff": diff}, open(p, "w"), indent=1, ensure_ascii=False)
        written.append(p)
    if "csv" in formats:
        p = os.path.join(outdir, "assets.csv")
        cols = ["host", "names", "group", "env", "status", "status_label", "closure", "closure_label", "confidence", "php", "ports", "cms",
                "server", "action", "note", "verified", "due", "overdue_days", "owner"]
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            for a in assets:
                w.writerow({k: safe_cell(v) for k, v in _asset_row(a, lang).items()})
        written.append(p)
    for fmt, name in (("nessus", "findings.nessus"), ("sarif", "assets.sarif"), ("defectdojo", "defectdojo.json"),
                      ("findings", "findings.json")):
        if fmt in formats:
            p = os.path.join(outdir, name)
            X.write(fmt, p, assets, summary, meta, lang)
            written.append(p)
    if "html" in formats:
        p = os.path.join(outdir, "report.html")
        open(p, "w", encoding="utf-8").write(html_report(assets, summary, groups, meta, diff, lang))
        written.append(p)
    if "xlsx" in formats:
        try:
            p = os.path.join(outdir, "report.xlsx")
            xlsx_report(p, assets, summary, groups, meta, diff, ep_rows, lang)
            written.append(p)
        except ImportError:
            print("[!] openpyxl not installed: skipping xlsx (pip install openpyxl)")
    if "docx" in formats:
        try:
            p = os.path.join(outdir, "report.docx")
            docx_report(p, assets, summary, groups, meta, diff, lang)
            written.append(p)
        except ImportError:
            print("[!] python-docx not installed: skipping docx (pip install python-docx)")
    return written
